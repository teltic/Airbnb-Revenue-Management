"""Preserve manual state across reruns: Notes on the main tabs and the
hand-maintained Promo Tracker tabs. Everything else in the workbook is
regenerated fresh every run.
"""

from __future__ import annotations

import datetime as dt
import logging
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

MAIN_TAB_NOTE_DATE_COL = 13  # M
MAIN_TAB_NOTE_COL = 14  # N
MAIN_TAB_DATE_COL = 1  # A
MAIN_TAB_HEADER_ROW = 1


def load_preserved_notes(path: Path, tab_name: str) -> dict[dt.date, tuple[object, str]]:
    """Returns {date: (note_date_value, note_text)} for a property's main tab,
    keyed by exact date match so a rerun can splice these back in even if
    row order or window length changed.
    """
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, data_only=False)
    if tab_name not in wb.sheetnames:
        return {}
    ws = wb[tab_name]
    preserved = {}
    for row in ws.iter_rows(min_row=MAIN_TAB_HEADER_ROW + 1):
        date_cell = row[MAIN_TAB_DATE_COL - 1]
        if not isinstance(date_cell.value, (dt.date, dt.datetime)):
            continue
        d = date_cell.value.date() if isinstance(date_cell.value, dt.datetime) else date_cell.value
        note_date = row[MAIN_TAB_NOTE_DATE_COL - 1].value
        note = row[MAIN_TAB_NOTE_COL - 1].value
        if note_date is not None or note is not None:
            preserved[d] = (note_date, note)
    return preserved


def load_promo_tab_rows(path: Path, promo_tab_name: str) -> list[dict[str, object]]:
    """Returns each row of a Promo Tracker tab as {header_name: value},
    keyed by whatever header text is actually in row 1 of the existing
    file -- not by column position. This makes reruns robust to the sheet
    schema changing between versions of this script (a column added,
    reordered, or removed): an older file's rows just won't have a key for
    a header that didn't exist yet, and callers use `.get()` accordingly,
    rather than needing an explicit migration step keyed to file version.
    """
    if not path.exists():
        return []
    wb = openpyxl.load_workbook(path, data_only=False)
    if promo_tab_name not in wb.sheetnames:
        return []
    ws = wb[promo_tab_name]
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append({headers[i]: v for i, v in enumerate(row) if i < len(headers) and headers[i]})
    return rows
