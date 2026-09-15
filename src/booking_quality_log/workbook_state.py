"""Read back a prior "Booking Quality Log" workbook's Reservation IDs and
hand-typed manual columns, so a rerun can:

1. tell whether any brand-new confirmed booking has shown up since that
   run (the daily skip-if-nothing-new check), and
2. carry every still-visible booking's manual notes forward untouched.

Both needs read the same data, so one function serves both.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from .workbook_build import HEADER_ROW, SHEET_NAME

RESERVATION_ID_COL = 26  # Z
MANUAL_COLS_START = 27  # AA
MANUAL_COLS_COUNT = 6


def load_prior_reservation_state(path: Path) -> dict[str, tuple]:
    """Returns {reservation_id: (6 manual-column values)} for every row in
    the prior file's "Booking Quality Log" sheet -- including rows where
    all 6 manual values are still blank, since the keys alone are what the
    daily skip check needs.
    """
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, data_only=False)
    if SHEET_NAME not in wb.sheetnames:
        return {}
    ws = wb[SHEET_NAME]

    result: dict[str, tuple] = {}
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        if len(row) <= RESERVATION_ID_COL - 1:
            continue
        reservation_id = row[RESERVATION_ID_COL - 1].value
        if not reservation_id:
            continue
        manual_values = tuple(
            row[MANUAL_COLS_START - 1 + i].value if len(row) > MANUAL_COLS_START - 1 + i else None
            for i in range(MANUAL_COLS_COUNT)
        )
        result[str(reservation_id)] = manual_values
    return result
