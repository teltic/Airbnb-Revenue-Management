import datetime as dt
from pathlib import Path

import openpyxl

from booking_quality_log.compute import BookingRow
from booking_quality_log.workbook_build import build_booking_quality_sheet, new_workbook
from booking_quality_log.workbook_state import load_prior_reservation_state


def _row(res_id):
    return BookingRow(
        property_name="Test Property",
        check_in=dt.date(2026, 9, 5),
        check_out=dt.date(2026, 9, 7),
        nights=2,
        stay_pattern="Weekend-anchored",
        one_night_stay=False,
        booked_date=dt.date(2026, 8, 20),
        booking_window_days=16,
        bw_vs_median="Typical",
        my_adr=400.0,
        my_revenue=800.0,
        source="airbnb",
        target_adr_p75=380.0,
        vs_target_dollar=20.0,
        vs_target_pct=0.0526,
        market_p25=250.0,
        market_p90=420.0,
        stly_adr="no data",
        demand_tier="High",
        gap_before_days="first known booking in window",
        gap_before_signal=None,
        gap_after_days=1,
        gap_after_signal="Upsell candidate (rarely fills alone)",
        status="Confirmed",
        reservation_id=res_id,
    )


def test_missing_file_returns_empty_dict(tmp_path):
    assert load_prior_reservation_state(tmp_path / "nope.xlsx") == {}


def test_reads_reservation_ids_and_manual_columns(tmp_path):
    wb = new_workbook()
    build_booking_quality_sheet(wb, [_row("AAA111"), _row("BBB222")], dt.date(2026, 9, 1))
    ws = wb["Booking Quality Log"]
    ws.cell(row=5, column=28, value="checked on airbnb")  # Comp Check for AAA111
    ws.cell(row=6, column=33, value="looks great")  # Notes/Verdict for BBB222
    path = tmp_path / "prior.xlsx"
    wb.save(path)

    state = load_prior_reservation_state(path)
    assert set(state.keys()) == {"AAA111", "BBB222"}
    assert state["AAA111"][0] == "checked on airbnb"
    assert state["BBB222"][5] == "looks great"


def test_row_with_all_blank_manual_columns_still_counted_for_skip_check(tmp_path):
    wb = new_workbook()
    build_booking_quality_sheet(wb, [_row("AAA111")], dt.date(2026, 9, 1))
    path = tmp_path / "prior.xlsx"
    wb.save(path)

    state = load_prior_reservation_state(path)
    assert "AAA111" in state
    assert state["AAA111"] == (None, None, None, None, None, None)


def test_reads_an_older_layout_without_a_status_column(tmp_path):
    # Simulates a file written by a version of this tool before the Status
    # column existed: "Reservation ID" sits at Z (26) instead of AA (27),
    # with manual columns starting right after it at AA (27). Locating
    # "Reservation ID" by its header text (not a hardcoded column number)
    # is what lets this old-layout file's notes still carry forward
    # correctly into a newer-layout file.
    wb = openpyxl.Workbook()
    del wb["Sheet"]
    ws = wb.create_sheet("Booking Quality Log")
    headers = [
        "Property", "Check-in", "In Day", "Check-out", "Out Day", "Nights",
        "Stay Pattern", "1-Night Stay", "Booked", "Booking Window (d)",
        "BW vs Median", "My ADR", "My Revenue", "Source", "Target ADR (P75)",
        "vs Target ($)", "vs Target (%)", "Market P25", "Market P90",
        "STLY ADR", "Demand Tier", "Gap Before (d)", "Gap Before Signal",
        "Gap After (d)", "Gap After Signal", "Reservation ID",
        "Comp Check (Airbnb)", "LY Weekday Occ.", "Pacing Push %",
        "LOS Discount", "Final PL Check", "Notes / Verdict",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=h)
    ws.cell(row=5, column=26, value="OLDID123")  # Z: Reservation ID (old position)
    ws.cell(row=5, column=27, value="checked, looked fine")  # old AA: Comp Check
    path = tmp_path / "old_layout.xlsx"
    wb.save(path)

    state = load_prior_reservation_state(path)
    assert "OLDID123" in state
    assert state["OLDID123"][0] == "checked, looked fine"
