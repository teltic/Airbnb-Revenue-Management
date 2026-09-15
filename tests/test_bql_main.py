import datetime as dt
from unittest.mock import MagicMock, patch

import openpyxl

from booking_quality_log import main as bql_main
from daily_price_analysis.config import Listing

FAKE_MARKET = {
    "Neighborhood Data Source": "Market Dashboard: test comp",
    "Future Occ/New/Canc": {
        "Category": {
            "test": {
                "Listings Used": 5,
                "X_values": ["2026-09-05", "2026-09-06"],
                "Y_values": [[70, 70], [0, 0], [0, 0], [40, 40], [40, 40], [0, 0]],
            }
        }
    },
    "Future Percentile Prices": {
        "Category": {
            "test": {
                "Listings Used": 5,
                "X_values": ["2026-09-05", "2026-09-06"],
                "Y_values": [[250, 250], [300, 300], [380, 380], [320, 320], [420, 420], [5, 5]],
            }
        }
    },
}

LISTINGS = [
    Listing(name="Test Property", tab_name="Test Property", short_suffix="Test", pms="testpms", listing_id="id-1")
]


def _reservation_row(res_id, checkin, checkout, confirmation_code, status="booked"):
    nights = (dt.date.fromisoformat(checkout) - dt.date.fromisoformat(checkin)).days
    return {
        "reservation_id": res_id,
        "listing_name": "Test Property",
        "check_in": checkin,
        "check_out": checkout,
        "booking_status": status,
        "rental_revenue": str(300 * nights),
        "no_of_days": nights,
        "booked_date": "2026-08-20",
        "booking_channel": "airbnb",
        "channelConfirmationCode": confirmation_code,
    }


def _make_client(reservations):
    client = MagicMock()
    client.get_reservations.return_value = reservations
    client.get_neighborhood_data.return_value = FAKE_MARKET
    return client


def _run(tmp_path, today, reservations):
    with patch("booking_quality_log.main.cfg.load_listings", return_value=LISTINGS), \
         patch("booking_quality_log.main.cfg.get_api_key", return_value="fake"), \
         patch("booking_quality_log.main.PriceLabsClient", return_value=_make_client(reservations)):
        return bql_main.run(tmp_path, today=today)


def test_first_run_always_writes(tmp_path):
    res = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    wrote = _run(tmp_path, dt.date(2026, 9, 15), res)
    assert wrote is True
    assert len(list(tmp_path.glob("*.xlsx"))) == 1


def test_no_new_booking_skips_even_after_a_multi_day_gap(tmp_path):
    res = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    _run(tmp_path, dt.date(2026, 9, 15), res)

    wrote = _run(tmp_path, dt.date(2026, 9, 18), res)  # same booking, 3 days later
    assert wrote is False
    assert len(list(tmp_path.glob("*.xlsx"))) == 1


def test_new_booking_writes_and_carries_manual_notes_forward(tmp_path):
    res_day1 = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    _run(tmp_path, dt.date(2026, 9, 15), res_day1)

    [day1_file] = tmp_path.glob("*.xlsx")
    wb = openpyxl.load_workbook(day1_file)
    ws = wb["Booking Quality Log"]
    ws.cell(row=5, column=27, value="checked, price is fair")
    wb.save(day1_file)

    res_day3 = res_day1 + [_reservation_row("r2", "2026-09-20", "2026-09-22", "BBBB222222")]
    wrote = _run(tmp_path, dt.date(2026, 9, 20), res_day3)
    assert wrote is True

    files = sorted(tmp_path.glob("*.xlsx"))
    assert len(files) == 2

    wb2 = openpyxl.load_workbook(files[-1])
    ws2 = wb2["Booking Quality Log"]
    rows = list(ws2.iter_rows(min_row=5, values_only=True))
    ids = {r[25] for r in rows}
    assert ids == {"AAAA111111", "BBBB222222"}

    carried = next(r for r in rows if r[25] == "AAAA111111")
    assert carried[26] == "checked, price is fair"


def test_cancellation_alone_does_not_trigger_a_new_file(tmp_path):
    res = [
        _reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111"),
        _reservation_row("r2", "2026-09-10", "2026-09-11", "BBBB222222"),
    ]
    _run(tmp_path, dt.date(2026, 9, 15), res)

    # r2 vanishes from the confirmed set (as if cancelled) -- no brand-new
    # booking appeared, so per Thomas's choice this must still skip.
    res_after_cancellation = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    wrote = _run(tmp_path, dt.date(2026, 9, 16), res_after_cancellation)
    assert wrote is False
    assert len(list(tmp_path.glob("*.xlsx"))) == 1


def test_force_writes_even_with_no_new_booking(tmp_path):
    res = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    _run(tmp_path, dt.date(2026, 9, 15), res)

    with patch("booking_quality_log.main.cfg.load_listings", return_value=LISTINGS), \
         patch("booking_quality_log.main.cfg.get_api_key", return_value="fake"), \
         patch("booking_quality_log.main.PriceLabsClient", return_value=_make_client(res)):
        wrote = bql_main.run(tmp_path, today=dt.date(2026, 9, 18), force=True)
    assert wrote is True
    assert len(list(tmp_path.glob("*.xlsx"))) == 2
